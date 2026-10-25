"""
Data processing utilities for the Client & Caregiver Streamlit dashboard.
This module reads the supplied Excel workbook and computes metrics from source data.
No source workbook values are modified.
"""
import io
import re
import pandas as pd

CLIENT_HINTS = {"customer id", "client name", "service type", "client pay"}
CAREGIVER_HINTS = {"name", "age", "present address", "education qualification", "work experience", "mobile"}

def normalize_text(value):
    if pd.isna(value):
        return value
    return re.sub(r"\s+", " ", str(value)).strip()

def _clean_headers(df):
    out = df.copy()
    out.columns = [normalize_text(c) if not pd.isna(c) else "" for c in out.columns]
    return out

def _read_sheets(file_bytes=None, path=None):
    """Read worksheets while expanding Excel merged cells into each covered row.

    Excel stores a merged range's value only in its top-left cell. The dashboard
    needs that value repeated for every record covered by the merged range, so
    expand merged ranges before converting the worksheet to a DataFrame.
    """
    import openpyxl
    source = io.BytesIO(file_bytes) if file_bytes is not None else path
    workbook = openpyxl.load_workbook(source, data_only=True, read_only=False)
    sheets = {}
    for ws in workbook.worksheets:
        max_row, max_col = ws.max_row, ws.max_column
        grid = [[ws.cell(r, c).value for c in range(1, max_col + 1)]
                for r in range(1, max_row + 1)]
        for merged_range in ws.merged_cells.ranges:
            min_col, min_row, max_col_range, max_row_range = merged_range.bounds
            value = ws.cell(min_row, min_col).value
            for row_i in range(min_row, max_row_range + 1):
                for col_i in range(min_col, max_col_range + 1):
                    grid[row_i - 1][col_i - 1] = value
        raw = pd.DataFrame(grid)
        if raw.empty:
            sheets[ws.title] = pd.DataFrame()
            continue
        # Find the first non-empty row most likely to contain headers.
        scan = min(len(raw), 8)
        scores = []
        for row_i in range(scan):
            vals = [normalize_text(x) if pd.notna(x) else "" for x in raw.iloc[row_i].tolist()]
            scores.append((sum(bool(x) for x in vals), row_i, vals))
        _, header_i, header_values = max(scores, key=lambda item: item[0]) if scores else (0, 0, [])
        if header_values and sum(bool(x) for x in header_values) >= 2:
            raw = raw.iloc[header_i + 1:].copy()
            unique_headers = []
            seen_headers = {}
            for col_i, header in enumerate(header_values):
                name = str(header).strip() if header else f"Unnamed: {col_i}"
                count = seen_headers.get(name, 0)
                seen_headers[name] = count + 1
                if count:
                    name = f"{name}_{count + 1}"
                unique_headers.append(name)
            raw.columns = unique_headers
            # Keep a separate, non-duplicated amount column for totals. When a
            # caregiver payment is merged across several client rows, show the
            # repeated amount on each row but count the original merged cell once.
            payment_header = next((h for h in unique_headers
                                   if str(h).strip().lower() in {
                                       "caregiver payment", "caregiver pay", "caregiver salary"
                                   }), None)
            if payment_header is not None:
                recorded = raw[payment_header].copy()
                payment_col_number = unique_headers.index(payment_header) + 1
                for merged_range in ws.merged_cells.ranges:
                    min_col, min_row, max_col_range, max_row_range = merged_range.bounds
                    if min_col == payment_col_number and max_col_range == payment_col_number and max_row_range > min_row:
                        for excel_row in range(min_row + 1, max_row_range + 1):
                            zero_based_index = excel_row - 1
                            if zero_based_index in recorded.index:
                                recorded.loc[zero_based_index] = None
                raw["Caregiver Payment Recorded"] = recorded
        else:
            raw.columns = [str(i) for i in range(raw.shape[1])]
        raw = _clean_headers(raw)
        raw = raw.dropna(how="all").reset_index(drop=True)
        drop = [c for c in raw.columns if (not str(c).strip() or str(c).lower().startswith("unnamed")) and raw[c].isna().all()]
        if drop:
            raw = raw.drop(columns=drop)
        sheets[ws.title] = raw
    workbook.close()
    return sheets

def _canonicalize_client(df):
    mapping = {}
    aliases = {
        "Sl No": ["sl no", "sl", "serial number", "serial no", "sl no."],
        "Customer ID": ["customer id", "client id", "customerid"],
        "Client Name": ["client name", "customer name", "name"],
        "Service Categories": ["service categories", "service category", "categories"],
        "Service Area": ["service area", "area", "service location"],
        "Service Type": ["service type", "type of service"],
        "Client Pay": ["client pay", "client payment", "customer pay"],
        "Caregiver Payment": ["caregiver payment", "caregiver pay", "caregiver salary"],
    }
    lookup = {str(c).strip().lower(): c for c in df.columns}
    for target, names in aliases.items():
        for name in names:
            if name in lookup:
                mapping[lookup[name]] = target
                break
    return df.rename(columns=mapping)

def _canonicalize_caregiver(df):
    mapping = {}
    aliases = {
        "SL": ["sl", "sl no", "serial number", "serial no"],
        "Name": ["name", "caregiver name", "full name"],
        "Age": ["age"],
        "Present Address": ["present address", "address", "location"],
        "Education Qualification": ["education qualification", "education", "qualification"],
        "Work Experience": ["work experience", "experience"],
        "Mobile": ["mobile", "mobile number", "phone", "phone number", "contact"],
    }
    lookup = {str(c).strip().lower(): c for c in df.columns}
    for target, names in aliases.items():
        for name in names:
            if name in lookup:
                mapping[lookup[name]] = target
                break
    return df.rename(columns=mapping)

def load_workbook_data(path=None, file_bytes=None):
    sheets = _read_sheets(file_bytes=file_bytes, path=path)
    client_name, caregiver_name = None, None
    client_score, caregiver_score = -1, -1
    for name, df in sheets.items():
        cols = {str(c).strip().lower() for c in df.columns}
        cs = len(cols & CLIENT_HINTS)
        gs = len(cols & CAREGIVER_HINTS)
        if cs > client_score:
            client_score, client_name = cs, name
        if gs > caregiver_score:
            caregiver_score, caregiver_name = gs, name
    # Prefer sheets whose actual headers identify the dataset.
    if client_score <= 0:
        client_name = list(sheets.keys())[0]
    if caregiver_score <= 0 or caregiver_name == client_name:
        remaining = [s for s in sheets if s != client_name]
        caregiver_name = remaining[0] if remaining else client_name
    clients = _canonicalize_client(sheets.get(client_name, pd.DataFrame())).copy()
    caregivers = _canonicalize_caregiver(sheets.get(caregiver_name, pd.DataFrame())).copy()
    # Preserve original workbook labels for diagnosis; canonical names are used by the app.
    metadata = {
        "sheet_rows": {name: int(len(df)) for name, df in sheets.items()},
        "client_sheet": client_name,
        "caregiver_sheet": caregiver_name,
        "original_columns": {name: [str(c) for c in df.columns] for name, df in sheets.items()},
    }
    return {"clients": clients, "caregivers": caregivers, "metadata": metadata}

def safe_numeric(series):
    if series is None:
        return series
    return pd.to_numeric(series.astype(str).str.replace(",", "", regex=False).str.replace("৳", "", regex=False).str.strip().replace({"": None, "nan": None, "None": None}), errors="coerce")

def completeness_report(df):
    if df.empty:
        return 0, 0
    complete = int(df.notna().all(axis=1).sum())
    return complete, int(len(df) - complete)

def format_money(value):
    if value is None or pd.isna(value):
        return "N/A"
    return f"{float(value):,.2f}".rstrip("0").rstrip(".")

def to_csv_bytes(df):
    return df.to_csv(index=False).encode("utf-8-sig")

def to_excel_bytes(data):
    if isinstance(data, pd.DataFrame):
        data = {"Data": data}
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, df in data.items():
            safe_name = str(name)[:31] or "Data"
            df.to_excel(writer, sheet_name=safe_name, index=False)
    return buffer.getvalue()

def quality_issues(clients, caregivers, metadata=None):
    rows = []
    def add_missing(df, dataset, id_col):
        for idx, row in df.iterrows():
            missing = [str(c) for c in df.columns if pd.isna(row[c]) or (isinstance(row[c], str) and not row[c].strip())]
            if missing:
                ident = row.get(id_col, idx + 1)
                rows.append({"Worksheet": dataset, "Record ID": ident, "Issue": "Missing information", "Details": "Blank fields: " + ", ".join(missing)})
    add_missing(clients, (metadata or {}).get("client_sheet", "Clients"), "Customer ID")
    add_missing(caregivers, (metadata or {}).get("caregiver_sheet", "Caregivers"), "SL")
    for df, dataset, key in [(clients,(metadata or {}).get("client_sheet","Clients"),"Customer ID"),
                             (caregivers,(metadata or {}).get("caregiver_sheet","Caregivers"),"SL")]:
        if key in df:
            dup = df[key].notna() & df[key].astype(str).duplicated(keep=False)
            for idx in df.index[dup]:
                rows.append({"Worksheet":dataset,"Record ID":df.loc[idx,key],"Issue":"Duplicate ID","Details":f"Repeated value in {key}"})
    for col in ["Client Pay","Caregiver Payment"]:
        if col in clients:
            numeric = safe_numeric(clients[col])
            invalid = clients[col].notna() & numeric.isna()
            for idx in clients.index[invalid]:
                rows.append({"Worksheet":(metadata or {}).get("client_sheet","Clients"),"Record ID":clients.loc[idx].get("Customer ID",idx+1),"Issue":"Invalid numeric value","Details":f"{col} is not numeric: {clients.loc[idx,col]}"})
    if not rows:
        return pd.DataFrame(columns=["Worksheet","Record ID","Issue","Details"])
    return pd.DataFrame(rows).drop_duplicates().reset_index(drop=True)

def make_summary(clients, caregivers, metadata=None):
    meta = metadata or {}
    items = [
        ("Generated from", "Loaded Excel workbook"),
        ("Client worksheet", meta.get("client_sheet", "Unknown")),
        ("Caregiver worksheet", meta.get("caregiver_sheet", "Unknown")),
        ("Total client records", len(clients)),
        ("Total caregiver records", len(caregivers)),
    ]
    if "Service Type" in clients:
        items.append(("Distinct service types", int(clients["Service Type"].dropna().astype(str).str.strip().replace("", pd.NA).dropna().nunique())))
    if "Service Area" in clients:
        items.append(("Distinct service areas", int(clients["Service Area"].dropna().astype(str).str.strip().replace("", pd.NA).dropna().nunique())))
    for col in ["Client Pay","Caregiver Payment"]:
        if col in clients:
            metric_col = "Caregiver Payment Recorded" if col == "Caregiver Payment" and "Caregiver Payment Recorded" in clients else col
            values = safe_numeric(clients[metric_col])
            items.append((f"Recorded total — {col}", values.sum(min_count=1)))
            source_values = safe_numeric(clients[col])
            items.append((f"Missing values — {col}", int(source_values.isna().sum())))
    return pd.DataFrame(items, columns=["Metric","Value"])
